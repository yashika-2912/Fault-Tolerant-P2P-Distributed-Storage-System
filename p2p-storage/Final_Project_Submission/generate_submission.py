from pathlib import Path
import csv,json,platform,sys,importlib.metadata as md,os
from datetime import datetime,timezone
ROOT=Path(__file__).resolve().parent
PRES=ROOT/'01_Final_Presentation'; DATA=ROOT/'03_Data_Associated_For_The_Paper'; MAN=ROOT/'04_Implementation_Manual'
for p in [PRES/'diagrams',DATA/'analysis',DATA/'plots',DATA/'raw_results',MAN/'diagrams']: p.mkdir(parents=True,exist_ok=True)
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DI,Pt as DP,RGBColor as DC
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

def csvout(name,headers,rows=[]):
 with (DATA/name).open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.DictWriter(f,fieldnames=headers);w.writeheader();w.writerows(rows)
inv=[
 ['SRC-TEST','tests/test_*.py','test source','Existing unit-test source modules, not measurements','Python','Repository','Existing source','6 modules','11 tests reported passing in prior audit'],
 ['BENCH','benchmark_results.csv','measurement table','Upload/download observations','CSV','benchmark harness not executed','None','0','No measurements'],
 ['FAULT','fault_tolerance_results.csv','measurement table','Failure/recovery observations','CSV','harness not executed','None','0','No measurements'],
 ['INTEGRITY','integrity_verification_results.csv','measurement table','Integrity test outcomes','CSV','tests/test_merkle.py and tests/test_merkle_audit.py; not rerun for package','None','0','No package-level outcomes'],
 ['REPLICA','replication_results.csv','measurement table','Observed replica placement outcomes','CSV','tests/test_placement.py; not rerun for package','None','0','No observed placements'],
 ['TEST-INPUT','test_file_manifest.csv','manifest','Persistent test-file inventory','CSV','No persistent test payload selected','None','0','Header only; no user files included'],
 ['ENV','environment_details.json','environment metadata','Captured runtime and configured node values','JSON','Project environment inspection','Environment query and source config','1 object','Available; hardware details incomplete']]
csvout('dataset_inventory.csv',['dataset_id','filename','dataset_type','description','format','source','generation_method','record_count','status'],[dict(zip(['dataset_id','filename','dataset_type','description','format','source','generation_method','record_count','status'],r)) for r in inv])
csvout('test_file_manifest.csv',['test_file_id','filename','file_type','file_size_bytes','checksum_algorithm','original_checksum','test_purpose','generation_method'])
csvout('benchmark_results.csv',['run_id','timestamp_utc','scenario','file_size_bytes','node_count','replication_factor','upload_time_ms','download_time_ms','throughput_mbps','trial_number','status','environment_id'])
csvout('fault_tolerance_results.csv',['run_id','failed_node','failure_type','detection_time_ms','promotion_time_ms','recovery_time_ms','file_retrieval_success','data_integrity_pass','replica_restoration_success','observations'])
csvout('integrity_verification_results.csv',['test_id','file_id','modification_type','expected_result','actual_result','integrity_check_method','detected_corruption','recovery_performed','status'])
csvout('replication_results.csv',['test_id','file_id','replication_factor_configured','replicas_expected','replicas_observed','node_failure_scenario','retrieval_success','consistency_status','notes'])
env={'captured_at_utc':datetime.now(timezone.utc).isoformat(),'os':platform.platform(),'python':'3.11.9 (project execution environment)','architecture':platform.machine(),'logical_cpu_count':os.cpu_count(),'memory_bytes':None,'hardware_note':'CPU model and memory not captured','packages':{'flask':'3.1.3','pycryptodome':'3.23.0','pytest':'9.1.1','matplotlib':'3.11.1','requests':'2.32.5'},'versions_note':'Inspected in the project execution environment; documentation-generation environment is separate.','nodes':{'N1':'ACTIVE','N2':'ACTIVE','N3':'STANDBY','replication_factor':2,'chunk_size_bytes':1048576,'heartbeat_interval_s':1,'heartbeat_timeout_s':4,'ports':[5000,6001,6002,6003,6100]},'experiments_executed':False}
(DATA/'environment_details.json').write_text(json.dumps(env,indent=2)+'\n',encoding='utf-8')
(DATA/'README.md').write_text("""# Experimental data package
No benchmark or end-to-end recovery experiment was run for this package. Result CSVs intentionally contain headers only; missing observations are not zeros or successes. Inventory documents sources and evidence status. Environment JSON records captured runtime/configuration values. Methodology and limitations describe future safe reproduction. No plots are generated without actual measured data.
Existing harnesses are experiments/bench_upload_retrieval.py and experiments/bench_failure_recovery.py; they start child processes on fixed localhost ports and write into experiments/results. They were not invoked to avoid port conflicts and source-tree output. `analysis/analyze_benchmarks.py` summarizes successful measured rows and generates a latency plot only if valid observations exist. Throughput, if computed, is decimal Mbps: file_size_bytes*8/(upload_time_ms/1000)/1,000,000. CSVs are UTF-8; durations use milliseconds and sizes use bytes.
""",encoding='utf-8')
(DATA/'methodology.md').write_text("""# Methodology and evidence status
No experiments were executed for this package. Existing harness defaults: upload/retrieval sizes 4096, 65536, 1048576 bytes, three trials; round-trip checks compare returned bytes. Failure harness uses 16384 bytes, terminates child N1, measures elapsed time to DOWN and N3 ACTIVE, then verifies download bytes. Before a controlled run, use an isolated clone, confirm ports 5000, 6001-6003, 6100 are free, preserve old results, record OS/hardware/package versions. Keep failures in raw records. For replication, inspect actual SQLite replica rows; do not infer observed placements from configured factor. Integrity detection must be reported separately from repair. Latencies are elapsed wall-clock milliseconds; throughput is decimal Mbps using original bytes and upload elapsed time.
""",encoding='utf-8')
(DATA/'limitations.md').write_text("""# Evidence limitations
No benchmark or end-to-end recovery measurements were collected. The prior 11-test result was repository unit testing, not a performance dataset. Fixed loopback ports mean scripts can conflict with active instances. The prototype has three locally simulated nodes and one coordinator/database; results would be host-specific. CPU model/RAM were not captured. Concurrent failures, partitions, and restart reintegration were not experimentally demonstrated. No numeric performance or recovery-success claim can be made.
""",encoding='utf-8')
# diagram images created via matplotlib
diags={}
def diagram(file,title,lines):
 fig,ax=plt.subplots(figsize=(12,4.5));ax.axis('off');ax.set_xlim(0,1);ax.set_ylim(0,1);ax.text(.03,.92,title,size=20,weight='bold',color='#1a365d')
 for i,(x,y,t) in enumerate(lines):
  ax.text(x,y,t,ha='center',va='center',size=12,bbox=dict(boxstyle='round,pad=.7',fc='#eaf3f7',ec='#009d9b',lw=2))
 for a,b,label in [(0,1,'HTTP'),(1,2,'metadata'),(1,3,'TCP'),(1,4,'TCP')]:
  if len(lines)>b:
   x1,y1,_=lines[a];x2,y2,_=lines[b];ax.annotate('',(x2,y2+.06),(x1,y1-.06),arrowprops=dict(arrowstyle='->',lw=1.8,color='#53718b'))
 fig.savefig(PRES/'diagrams'/f'{file}.png',dpi=140,bbox_inches='tight');fig.savefig(PRES/'diagrams'/f'{file}.svg',bbox_inches='tight');plt.close(fig)
 (MAN/'diagrams'/f'{file}.png').write_bytes((PRES/'diagrams'/f'{file}.png').read_bytes())
 return PRES/'diagrams'/f'{file}.png'
diags['architecture']=diagram('architecture','Verified local prototype architecture',[(.1,.55,'CLI / HTTP'),(.36,.55,'Flask coordinator'),(.64,.55,'SQLite metadata'),(.88,.78,'N1 + N2 ACTIVE'),(.88,.28,'N3 STANDBY')])
diags['upload']=diagram('upload_flow','Upload: request → split → encrypt → replicate → metadata',[(.12,.5,'POST /files'),(.36,.5,'1 MiB chunks'),(.61,.5,'AES-256-GCM'),(.84,.5,'STORE to 2 active nodes'),(.61,.2,'SQLite after ACK')])
diags['download']=diagram('download_flow','Download: fetch → per-chunk check/decrypt → root → reconstruct',[(.12,.5,'GET /files/id'),(.36,.5,'Fetch replica'),(.61,.5,'SHA-256 check'),(.84,.5,'AES-GCM decrypt'),(.61,.2,'Merkle root → reassemble')])
diags['failure']=diagram('failure_recovery','Failure detection and standby repair',[(.12,.5,'Heartbeat 1 s'),(.36,.5,'Timeout >4 s'),(.61,.5,'Mark ACTIVE DOWN'),(.84,.5,'Attempt replica copy'),(.61,.2,'Promote standby')])
diags['upload_flow']=diags['upload'];diags['download_flow']=diags['download'];diags['failure_recovery']=diags['failure']
(MAN/'diagrams'/'README.md').write_text('See PNG diagrams for architecture and upload, retrieval, and failure flow. They describe inspected code paths.\n',encoding='utf-8')
# data docs
(DATA/'analysis'/'README.md').write_text('Run `python analysis/analyze_benchmarks.py` after adding genuine successful rows to benchmark_results.csv. It writes benchmark_summary.csv and generates a plot only when usable observations exist. Currently there are no usable rows, so no plot is emitted.\n',encoding='utf-8')
(DATA/'raw_results'/'README.md').write_text('No raw experimental data collected.\n',encoding='utf-8')
(PRES/'diagrams'/'README.md').write_text('Architecture and workflow images generated from repository-verified behavior. No experimental chart is included because no results were measured.\n',encoding='utf-8')

chapters=[
('Project overview','Educational local distributed storage prototype. Flask coordinates encrypted chunk uploads/downloads, SQLite tracks metadata, and three TCP nodes hold data. N1/N2 start ACTIVE; N3 starts STANDBY.'),
('Objectives and scope','Store encrypted chunks, retain two replicas, reconstruct downloads with integrity checks, and detect stale nodes and attempt standby repair. Source implements these mechanisms; package evidence does not include end-to-end measurement.'),
('System architecture','CLI uses HTTP with coordinator/app.py. Storage logic is coordinator/storage.py, metadata coordinator/db.py, nodes node/node_process.py. TCP node requests and heartbeat messages use length-prefixed JSON. HTTP port 5000, heartbeat 6100, nodes 6001-6003 on loopback.'),
('Directory and module guide','client/cli_client.py: HTTP client. coordinator/app.py: API and startup. storage.py: upload/download pipeline. chunking.py: 1 MiB chunks. crypto.py: AES-GCM/SHA-256. placement.py: deterministic pair. node_client.py: TCP STORE/FETCH. node/node_process.py: storage server/heartbeats. heartbeat_monitor.py: health detection. recovery.py: standby repair. db.py: schema. merkle.py: Merkle root. audit.py: hash-chain. tests/: component tests. experiments/: benchmark harnesses.'),
('Database and metadata','SQLite tables: files(file_id PK, filename, size_bytes, chunk_count, merkle_root, created_at); chunks(chunk_id PK, file_id FK, chunk_index, chunk_hash); nodes(node_id PK, host, port, status, last_heartbeat); replicas(chunk_id,node_id composite PK, role); audit_log(entry_index PK,event,prev_hash,entry_hash,created_at). Foreign keys enabled. P2P_DB_PATH overrides default coordinator/metadata.db.'),
('Upload workflow','POST /files multipart field file. Requires two ACTIVE nodes; split plaintext into 1 MiB chunks, AES-256-GCM encrypt each with random 12-byte nonce/16-byte tag, hash encrypted data, calculate Merkle root, place on two nodes via STORE, and persist metadata after all acknowledgements. Failed partial writes may leave orphan node files.'),
('Retrieval workflow','GET /files/<id> loads ordered chunks and replicas. Try nodes not marked DOWN; FETCH, compare encrypted SHA-256, decrypt/authenticate AES-GCM, verify Merkle root, reassemble, and verify original byte count. Missing usable replica causes error.'),
('Replication placement','SHA-256(chunk ID) first eight bytes big-endian modulo sorted ACTIVE node count chooses primary; next circular position is secondary. Initial two active nodes both receive each chunk. Fewer than two ACTIVE nodes prevents upload.'),
('Failure and repair','Heartbeats every 1 s; ACTIVE node older than 4 s after first heartbeat is marked DOWN. Recovery selects a STANDBY, fetches affected data from a surviving non-DOWN replica, stores on standby, remaps only after ACK, then promotes standby. No standby means no repair. Restarted DOWN nodes are not automatically reintegrated.'),
('Security and integrity','AES-256-GCM blob is nonce(12)||ciphertext||tag(16). Static demo key in coordinator/config.py; no external key management. HTTP/TCP lack authentication/TLS. SHA-256 checks encrypted chunks; Merkle combines digest strings. GCM authenticates ciphertext. Audit entries link hashes but are not externally immutable.'),
('Configuration','config.py: heartbeat 1 s, timeout 4 s, port 6100, chunk size 1 MiB, factor 2, static key, node ports 6001-6003, DB path coordinator/metadata.db or P2P_DB_PATH. HTTP port 5000 in app.py. node CLI requires --node-id and --port; coordinator host/port optional defaults loopback/6100.'),
('Installation','Python 3.10+ per README. From repo root: python -m venv .venv; PowerShell .venv\\Scripts\\Activate.ps1; python -m pip install -r requirements.txt. Dependencies are Flask, PyCryptodome, pytest, Matplotlib, Requests.'),
('Run commands','Four terminals: python coordinator/app.py; python node/node_process.py --node-id N1 --port 6001; repeat N2/6002 and N3/6003. Wait for heartbeats. CLI: python client/cli_client.py status; upload path; download FILE_ID output. Ctrl+C to stop.'),
('HTTP and TCP protocol','HTTP routes: POST /files, GET /files/<file_id>, GET /status. TCP messages use 4-byte big-endian length and UTF-8 JSON. STORE sends chunk id and base64 bytes; FETCH sends id. One request per connection. Heartbeat is a separate TCP listener.'),
('Tests and verification','Prior audit reported `pytest -q`: 11 passed; not rerun during package generation. Reproduce from repository root with `pytest -q` after installing requirements. Existing tests cover chunking, crypto, placement/schema, Merkle/audit, heartbeat detection. They do not establish benchmark performance or end-to-end recovery. Benchmark commands are `python experiments/bench_upload_retrieval.py` and `python experiments/bench_failure_recovery.py`; use an isolated clone, ensure fixed ports are free, and preserve existing results before execution.'),
('Performance evaluation','No numeric performance results. Existing upload benchmark defaults to 4096/65536/1048576 bytes, 3 trials, validates equality. Failure benchmark uses 16384 bytes and measures DOWN and N3 ACTIVE transitions over 3 trials. Neither was run because scripts use fixed ports and write into source tree.'),
('Troubleshooting','Dependencies: activate venv/install requirements. Port bind: inspect ports before use and do not stop unknown processes. Node status: check coordinator and node host/port/heartbeats. Upload 503: requires two ACTIVE nodes. Download failure: check replica records and nodes. Hash mismatch skips replica; repair is not automatic for corruption. Keep P2P_DB_PATH consistent.'),
('Limitations and future work','Single coordinator/SQLite, loopback defaults, no discovery/consensus/auth/TLS/key management/garbage collection/multi-coordinator HA/partition policy. One-node repair depends on surviving source and standby. No automatic reintegration. Future: secure keys, authenticated transport, resilient metadata, multi-host tests.'),
('Conclusion','Source implements encrypted replicated chunk storage and basic heartbeat-triggered standby repair. Existing unit tests were reported passing. No performance or end-to-end recovery outcome is claimed here.'),
('References','Flask docs: https://flask.palletsprojects.com/ ; SQLite: https://www.sqlite.org/docs.html ; Python sockets: https://docs.python.org/3/library/socket.html ; PyCryptodome AES: https://pycryptodome.readthedocs.io/en/latest/src/cipher/aes.html ; NIST FIPS 180-4: https://doi.org/10.6028/NIST.FIPS.180-4 ; pytest: https://docs.pytest.org/ ; Matplotlib: https://matplotlib.org/stable/')]
# Manual DOCX
doc=Document();sec=doc.sections[0];sec.top_margin=DI(.7);sec.bottom_margin=DI(.7);sec.left_margin=DI(.8);sec.right_margin=DI(.8)
doc.styles['Normal'].font.name='Aptos';doc.styles['Normal'].font.size=DP(10)
doc.styles['Heading 1'].font.name='Aptos Display';doc.styles['Heading 1'].font.size=DP(18);doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 1'].font.color.rgb=DC(26,54,93)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('Fault Tolerant P2P Distributed Storage System  |  Page ')
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('FAULT TOLERANT P2P\nDISTRIBUTED STORAGE SYSTEM');r.bold=True;r.font.size=DP(28);r.font.color.rgb=DC(26,54,93)
p=doc.add_paragraph('Implementation Manual');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Student: [Student Name]\nRoll: [Roll Number]\nInstitution: [Institution]\nDepartment: [Department]\nGuide: [Guide Name]\nAcademic year: [Academic Year]').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break();doc.add_heading('Document index',0)
for title,_ in chapters:doc.add_paragraph(title,style='List Number')
doc.add_paragraph('Evidence note: existing unit tests were reported as 11 passing; benchmark and end-to-end fault recovery remain unexecuted.')
for i,(title,body) in enumerate(chapters):
 doc.add_page_break();doc.add_heading(f'{i+1}. {title}',1);doc.add_paragraph(body)
 key={2:'architecture',5:'upload_flow',6:'download_flow',8:'failure_recovery'}.get(i)
 if key:
  doc.add_picture(str(diags[key]),width=DI(6.2));doc.add_paragraph(f'Figure {i+1}. Code-grounded workflow.')
docx=MAN/'Implementation_Manual.docx';doc.save(docx)
# PDF
ss=getSampleStyleSheet();ss.add(ParagraphStyle(name='ChapterX',parent=ss['Heading1'],fontSize=16,leading=20,textColor=colors.HexColor('#1a365d'),spaceAfter=12));ss.add(ParagraphStyle(name='BodyX',parent=ss['BodyText'],fontSize=9.2,leading=13,spaceAfter=8))
story=[Spacer(1,1.2*inch),Paragraph('<b>FAULT TOLERANT P2P<br/>DISTRIBUTED STORAGE SYSTEM</b>',ss['Title']),Paragraph('Implementation Manual',ss['Heading2']),Paragraph('Student: [Student Name]<br/>Roll: [Roll Number]<br/>Institution: [Institution]<br/>Department: [Department]<br/>Guide: [Guide Name]<br/>Academic year: [Academic Year]',ss['BodyText']),PageBreak(),Paragraph('Document index',ss['ChapterX'])]
story += [Paragraph(f'{i+1}. {x}',ss['BodyX']) for i,(x,_) in enumerate(chapters)]
story += [Paragraph('Evidence note: prior audit reports 11 unit tests passing. Benchmarks and end-to-end failure recovery were not executed.',ss['BodyX'])]
for i,(title,body) in enumerate(chapters):
 story += [PageBreak(),Paragraph(f'{i+1}. {title}',ss['ChapterX']),Paragraph(body.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'),ss['BodyX'])]
 key={2:'architecture',5:'upload_flow',6:'download_flow',8:'failure_recovery'}.get(i)
 if key:story += [Image(str(diags[key]),width=6.2*inch,height=2.4*inch),Paragraph(f'Figure: {title}',ss['BodyX'])]
def pdf_footer(canvas, document):
 canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#5a6c80'));canvas.drawString(.65*inch,.4*inch,'Fault Tolerant P2P Distributed Storage System | Implementation Manual');canvas.drawRightString(7.85*inch,.4*inch,str(document.page));canvas.restoreState()
pdf=MAN/'Implementation_Manual.pdf';SimpleDocTemplate(str(pdf),pagesize=letter).build(story,onFirstPage=pdf_footer,onLaterPages=pdf_footer)
(MAN/'README.md').write_text('# Implementation Manual\n\nOpen Implementation_Manual.pdf or edit Implementation_Manual.docx. The versions contain the same 20 chapters. Personal and institution fields remain placeholders. No measurement claims are made.\n',encoding='utf-8')
(MAN/'configuration_reference.md').write_text('# Configuration reference\n\n| Name | Value/source |\n|---|---|\n| HEARTBEAT_INTERVAL | 1 s, config.py |\n| HEARTBEAT_TIMEOUT | 4 s, config.py |\n| COORDINATOR_PORT | 6100, config.py |\n| HTTP port | 5000, app.py |\n| CHUNK_SIZE | 1 MiB |\n| REPLICATION_FACTOR | 2 |\n| AES_KEY | Static demo key in source; do not expose/use for sensitive data |\n| NODE_PORTS | 6001, 6002, 6003 |\n| DB_PATH | coordinator/metadata.db; P2P_DB_PATH override |\n| Node CLI | --node-id, --port required; --coordinator-host/port optional |\n\nThere are no environment overrides for other constants.\n',encoding='utf-8')
(MAN/'troubleshooting.md').write_text('# Troubleshooting\n\n- Install requirements in activated virtualenv for import errors.\n- Inspect port owners before startup; do not terminate unknown processes.\n- Node missing: verify coordinator listener and node host/port, inspect GET /status.\n- Upload 503: requires at least two ACTIVE nodes and successful STORE acknowledgments.\n- Download error: verify replica metadata and node availability.\n- AES authentication failure means bytes/key mismatch; do not treat SHA-256 as repair.\n- Keep P2P_DB_PATH consistent across coordinator and isolated processes.\n- Restarted DOWN nodes are not automatically reintegrated.\n',encoding='utf-8')
# Presentation data and speaker notes
deck=[
('Title',['Student: [Student Name] • Roll: [Roll Number]','Department/Institution: [Department] / [Institution]','Guide: [Guide Name] • Academic year: [Academic Year]']),
('Outline',['Problem and objectives','Verified architecture and workflows','Algorithms, security, and implementation','Testing evidence and experimental status','Limitations and conclusion']),
('Introduction',['Distributed storage spreads data across networked nodes.','Replication can preserve access to another copy after a node loss.','This repository is a local educational prototype.']),
('Problem statement',['A failed storage node can hide a sole copy.','Replica placement and metadata are needed to find chunks.','Retrieval should detect missing or corrupted data.','Repair requires an available copy and destination.']),
('Objectives',['Encrypt and split file contents into fixed chunks.','Store two replicas across active nodes.','Reconstruct data and validate stored chunks.','Detect stale nodes and attempt standby repair.']),
('Existing systems',['Centralized platforms can also use redundancy.','This project demonstrates mechanisms, not production availability.','A single coordinator and SQLite remain common failure points.','No consensus, partition policy, or multi-coordinator HA.']),
('Proposed system',['Flask coordinator manages API workflows.','SQLite persists file/chunk/node/replica/audit metadata.','N1/N2 ACTIVE; N3 STANDBY.','TCP nodes store local encrypted chunk blobs.']),
('Architecture',['Verified architecture is shown in the diagram.']),
('Upload workflow',['POST /files → 1 MiB split → AES-256-GCM.','Two active nodes receive each encrypted chunk.','Metadata written after all STORE acknowledgments.']),
('Retrieval workflow',['Read chunk and replica metadata.','Try non-DOWN replicas; verify SHA-256.','AES-GCM decrypt, verify Merkle root, reassemble and size-check.']),
('Methodology',['Bytes → chunks → encryption → placement → TCP STORE.','Persist metadata only after successful replica ACKs.','Download reverses path and checks integrity.']),
('Major modules',['app.py/API; storage.py pipeline; db.py metadata.','crypto.py/chunking.py/merkle.py/audit.py.','placement.py/node_client.py/node_process.py.','heartbeat_monitor.py and recovery.py.']),
('Algorithms',['Placement: SHA-256 chunk ID prefix modulo sorted active count; next circular node.','Failure: heartbeat age exceeds 4 s after first heartbeat.','Integrity: encrypted chunk SHA-256, Merkle root, GCM authentication.','Audit chain links event hash to prior event.']),
('Failure and repair',['Heartbeat → stale check → DOWN → copy affected chunks → promote standby.','Promotion can occur even when some individual repairs were skipped.','Automatic reintegration of a restarted DOWN node is absent.']),
('Security',['AES-256-GCM: random 12-byte nonce + ciphertext + 16-byte tag.','Static demo key; no external key management.','No HTTP/TCP authentication or TLS.','Hash is integrity evidence, not confidentiality.']),
('Implementation/tools',['Python, Flask, SQLite, TCP sockets, PyCryptodome.','Requests client, pytest, Matplotlib.','Versions appear in environment_details.json.']),
('Experimental setup',['3 nodes, replication factor 2; chunks 1 MiB.','Heartbeat interval 1 s, threshold 4 s.','Loopback ports: 5000, 6100, 6001–6003.','Hardware model and memory not recorded.']),
('Evaluation methods',['Upload/download harness defaults: 4 KiB, 64 KiB, 1 MiB; 3 trials.','Failure harness: 16 KiB; terminate N1; track DOWN and N3 ACTIVE.','Throughput formula documented in methodology.','Experiments unexecuted; fixed ports require isolation.']),
('Performance results',['No valid benchmark measurements are available.','CSV schema is present and contains no data rows.','No performance chart is shown to avoid invented results.']),
('Fault/integrity results',['Prior audit: 11 unit tests passed.','Tests cover components; no end-to-end recovery run.','No measured success rate or repair time is available.']),
('Discussion',['Implementation exists for placement, detection, repair, integrity.','Unit-test report does not establish runtime performance.','No automatic reintegration after restart.','One-node recovery depends on surviving replica and standby.']),
('Limitations/future',['Single coordinator/database and loopback configuration.','No secure key management, TLS, auth, peer discovery, garbage collection.','Concurrent failures and partitions unverified.','Future: HA metadata, secure transport, multi-host evaluation.']),
('Conclusion',['Encrypted two-replica chunk path and basic repair logic exist in code.','Prior unit tests reported passing.','Benchmarks and end-to-end recovery remain unmeasured.']),
('References / Questions',['Flask; SQLite; Python sockets; PyCryptodome AES.','NIST FIPS 180-4 SHA-2.','pytest and Matplotlib official documentation.','Questions'])]

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
BLUE=RGBColor(26,54,93);TEAL=RGBColor(0,157,155);INK=RGBColor(36,49,66);MUTED=RGBColor(90,108,128)
for i,(title,bullets) in enumerate(deck):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
 bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.14));bar.fill.solid();bar.fill.fore_color.rgb=TEAL;bar.line.fill.background()
 t=s.shapes.add_textbox(Inches(.65),Inches(.4),Inches(12),Inches(.6)).text_frame.paragraphs[0];t.text=title;t.font.name='Aptos Display';t.font.size=Pt(27);t.font.bold=True;t.font.color.rgb=BLUE
 if i==0:
  tb=s.shapes.add_textbox(Inches(.9),Inches(1.55),Inches(11.5),Inches(1.6));p=tb.text_frame.paragraphs[0];p.text='FAULT TOLERANT P2P\nDISTRIBUTED STORAGE SYSTEM';p.font.size=Pt(34);p.font.bold=True;p.font.color.rgb=BLUE
  y=3.7
 else:y=1.55
 if title=='Architecture':s.shapes.add_picture(str(diags['architecture']),Inches(.8),Inches(1.35),width=Inches(11.8),height=Inches(5.3))
 elif title=='Upload workflow':s.shapes.add_picture(str(diags['upload_flow']),Inches(.8),Inches(1.35),width=Inches(11.8),height=Inches(5.3))
 elif title=='Retrieval workflow':s.shapes.add_picture(str(diags['download_flow']),Inches(.8),Inches(1.35),width=Inches(11.8),height=Inches(5.3))
 elif title=='Failure and repair':s.shapes.add_picture(str(diags['failure_recovery']),Inches(.8),Inches(1.35),width=Inches(11.8),height=Inches(5.3))
 else:
  for b in bullets:
   sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(.95),Inches(y+.1),Inches(.15),Inches(.15));sh.fill.solid();sh.fill.fore_color.rgb=TEAL;sh.line.fill.background()
   tb=s.shapes.add_textbox(Inches(1.3),Inches(y),Inches(10.9),Inches(.72));p=tb.text_frame.paragraphs[0];p.text=b;p.font.name='Aptos';p.font.size=Pt(18);p.font.color.rgb=INK;y+=.95
 f=s.shapes.add_textbox(Inches(.65),Inches(7.05),Inches(10),Inches(.2)).text_frame.paragraphs[0];f.text='Fault Tolerant P2P Distributed Storage System';f.font.size=Pt(9);f.font.color.rgb=MUTED
 n=s.shapes.add_textbox(Inches(12.2),Inches(7.02),Inches(.5),Inches(.25)).text_frame.paragraphs[0];n.text=str(i+1);n.font.size=Pt(10);n.font.bold=True;n.font.color.rgb=TEAL
pptx=PRES/'Fault_Tolerant_P2P_Distributed_Storage_System.pptx';prs.save(pptx)
(PRES/'presentation_notes.md').write_text('# Presentation notes\n\n24-slide deck; explain diagrams as repository-grounded, and state that no experiment numbers are available.\n\n'+'\n'.join(f'## Slide {i+1}: {t}\n\n'+' '.join(b)+'\n' for i,(t,b) in enumerate(deck))+'\n## Viva terms\n\n- Replica: second stored copy.\n- Heartbeat: periodic liveness message.\n- Merkle root: aggregate digest over leaves.\n- AES-GCM: authenticated encryption.\n\n## Viva questions\n\n- Does hashing encrypt? No.\n- Is recovery guaranteed for multiple failures? No.\n- Are benchmarks available? No.\n',encoding='utf-8')
# Index
(ROOT/'SUBMISSION_INDEX.md').write_text('''# Final Project Submission Index

Project: Fault Tolerant P2P Distributed Storage System

- 01_Final_Presentation/Fault_Tolerant_P2P_Distributed_Storage_System.pptx — 24-slide widescreen deck.
- 01_Final_Presentation/presentation_notes.md — slide notes and viva terms.
- 01_Final_Presentation/diagrams/ — architecture and workflow PNG/SVG.
- 03_Data_Associated_For_The_Paper/ — CSV schemas, inventory, manifest, JSON environment, methodology, limitations, raw-results folder, `analysis/analyze_benchmarks.py`, and header-only `analysis/benchmark_summary.csv`. No fabricated measurement rows or plots.
- 04_Implementation_Manual/Implementation_Manual.docx and .pdf — equivalent 20-chapter manuals.
- Manual README, configuration reference, troubleshooting, diagrams.

Prior repository audit reported 11 unit tests passing. They were not rerun here. Benchmark and end-to-end fault-recovery experiments remain unexecuted. No plots are appropriate without measured data. Personal fields are placeholders. Static demo key, missing auth/TLS, single coordinator/SQLite, and no automatic node reintegration are known limits.

`validation_report.json` records automated structural checks. `generate_submission.py` is the reproducible artifact generator; its isolated tool environment is stored in `.docenv/`.
''',encoding='utf-8')
# validation
from pptx import Presentation as RPR
from docx import Document as RDO
r=RPR(str(pptx));d=RDO(str(docx));pdfbytes=pdf.read_bytes()
from pypdf import PdfReader
pdfreader=PdfReader(str(pdf));slide_text=[sh.text for sl in r.slides for sh in sl.shapes if getattr(sh,'has_text_frame',False)]
doc_text='\n'.join(x.text for x in d.paragraphs);pdf_text='\n'.join(x.extract_text() or '' for x in pdfreader.pages)
csv_rows={x.name:sum(1 for _ in csv.reader(x.open(encoding='utf-8-sig')))-1 for x in DATA.glob('*.csv')}
validation={'pptx_exists':pptx.exists(),'slide_count':len(r.slides),'slide_titles_count':len(deck),'slide_text_shape_count':len(slide_text),'pptx_text_present':all(title in '\n'.join(slide_text) for title,_ in deck),'docx_exists':docx.exists(),'docx_chapter_heading_count':sum(x.style.name=='Heading 1' for x in d.paragraphs),'docx_expected_chapters':len(chapters),'pdf_exists':pdf.exists(),'pdf_pages':len(pdfreader.pages),'pdf_header_valid':pdfbytes.startswith(b'%PDF-'),'pdf_eof_valid':pdfbytes.rstrip().endswith(b'%%EOF'),'pdf_has_chapter_text':all(title in pdf_text for title,_ in chapters),'csv_observation_rows':csv_rows,'json_valid':bool(json.loads((DATA/'environment_details.json').read_text())),'diagram_pngs':len(list(PRES.joinpath('diagrams').glob('*.png'))),'manual_diagram_pngs':len(list(MAN.joinpath('diagrams').glob('*.png'))),'analysis_script_exists':(DATA/'analysis'/'analyze_benchmarks.py').exists(),'analysis_summary_exists':(DATA/'analysis'/'benchmark_summary.csv').exists(),'measurement_rows':0,'benchmark_executed':False,'end_to_end_recovery_executed':False}
(ROOT/'validation_report.json').write_text(json.dumps(validation,indent=2)+'\n',encoding='utf-8')
print(json.dumps(validation,indent=2))

